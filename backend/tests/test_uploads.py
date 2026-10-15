import io
import json
import os
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

from fastapi import HTTPException
from fastapi.testclient import TestClient
from pypdf import PdfReader, PdfWriter

from app.api.v1 import documents
from app.config import settings
from app.core import local_db
from app.core.supabase_client import SafeSupabaseWrapper
from app.core.uploads import safe_filename, storage_path_from, validate_pdf_bytes
from app.services import document_pipeline
from main import app

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "fixtures"))
import make_fixtures  # noqa: E402

OTHER_USER = "bbbbbbbb-bbbb-bbbb-bbbb-bbbbbbbbbbbb"


class SafeFilenameTests(unittest.TestCase):
    def test_safe_filename_strips_traversal(self):
        self.assertEqual(safe_filename("../../etc/passwd.pdf"), "passwd.pdf")
        self.assertEqual(safe_filename("..\\..\\windows\\evil name.pdf"), "evil_name.pdf")
        self.assertEqual(safe_filename("Sách tiếng Việt (bản 2).pdf"), "S_ch_ti_ng_Vi_t_b_n_2.pdf")
        for name in ("", None, "..", "/", ".pdf", "???"):
            self.assertEqual(safe_filename(name), "document.pdf")

    def test_safe_filename_limits_length_and_keeps_extension(self):
        result = safe_filename("a" * 300 + ".pdf")
        self.assertEqual(len(result), 120)
        self.assertTrue(result.endswith(".pdf"))

    def test_storage_path_from_handles_legacy_public_urls(self):
        path = "user-1/doc-1/My Book.pdf"
        self.assertEqual(storage_path_from(path), path)
        self.assertEqual(
            storage_path_from("https://x.supabase.co/storage/v1/object/public/documents/user-1/doc-1/My%20Book.pdf?t=1"),
            path,
        )
        self.assertEqual(
            storage_path_from("http://localhost:8000/static/storage/documents/user-1/doc-1/My%20Book.pdf"), path
        )
        self.assertIsNone(storage_path_from("https://example.com/somewhere/else.pdf"))
        self.assertIsNone(storage_path_from(None))


class LocalStoreTestCase(unittest.TestCase):
    """Chạy trên LocalStore trong thư mục tạm, không đụng Supabase hay backend/data."""

    @classmethod
    def setUpClass(cls):
        cls.fixtures_dir = tempfile.TemporaryDirectory()
        cls.expected = make_fixtures.build_all(cls.fixtures_dir.name)
        with open(os.path.join(cls.fixtures_dir.name, "text_book.pdf"), "rb") as f:
            cls.pdf_bytes = f.read()

    @classmethod
    def tearDownClass(cls):
        cls.fixtures_dir.cleanup()

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        db_file = os.path.join(tmp.name, "local_db.json")
        with open(db_file, "w", encoding="utf-8") as f:
            f.write("{}")
        self.store = SafeSupabaseWrapper(None)
        for p in (
            patch.object(local_db, "DB_FILE", db_file),
            patch.object(local_db, "STORAGE_DIR", os.path.join(tmp.name, "storage")),
            patch.object(settings, "AUTH_REQUIRED", False),
            patch.object(settings, "ENVIRONMENT", "development"),
            patch.object(settings, "BACKEND_PUBLIC_URL", "http://testserver"),
            patch("app.api.v1.documents.get_supabase_client", return_value=self.store),
            patch("app.core.supabase_client.get_supabase_client", return_value=self.store),
            patch("app.api.v1.ebook.get_supabase_client", return_value=self.store),
            # Không để test gọi Gemini thật qua OCR
            patch("app.services.document_pipeline.get_ocr_provider", return_value=None),
            # Các test upload nhiều sách: không để giới hạn lượt miễn phí chen vào
            patch("app.services.billing.free_book_limit", return_value=100),
        ):
            p.start()
            self.addCleanup(p.stop)
        self.storage_dir = os.path.join(tmp.name, "storage")
        self.client = TestClient(app)

    def rows(self, table):
        return self.store.table(table).select("*").execute().data

    def upload(self, content, filename="book.pdf"):
        return self.client.post(
            "/api/v1/documents/upload-pdf",
            files={"file": (filename, content, "application/pdf")},
            data={"title": "Test book"},
        )


class UploadValidationTests(LocalStoreTestCase):
    def setUp(self):
        super().setUp()
        self.background = Mock()
        p = patch.object(documents, "process_uploaded_file", self.background)
        p.start()
        self.addCleanup(p.stop)

    def test_upload_rejects_non_pdf_magic(self):
        res = self.upload(b"<html>this is not a pdf</html>")
        self.assertEqual(res.status_code, 400)
        self.assertEqual(self.rows("documents"), [])
        self.background.assert_not_called()

    def test_upload_rejects_wrong_extension_and_empty_file(self):
        self.assertEqual(self.upload(self.pdf_bytes, filename="book.exe").status_code, 400)
        self.assertEqual(self.upload(b"").status_code, 400)

    def test_upload_rejects_oversize(self):
        with patch.object(settings, "MAX_PDF_MB", 1):
            res = self.upload(self.pdf_bytes + b"\0" * (1024 * 1024))
        self.assertEqual(res.status_code, 413)
        self.assertEqual(self.rows("documents"), [])

    def test_upload_rejects_too_many_pages(self):
        with patch.object(settings, "MAX_PDF_PAGES", 3):
            self.assertEqual(self.upload(self.pdf_bytes).status_code, 413)

    def test_upload_rejects_corrupt_and_password_protected_pdf(self):
        self.assertEqual(self.upload(b"%PDF-1.7\nnot really a pdf").status_code, 400)

        writer = PdfWriter(clone_from=PdfReader(io.BytesIO(self.pdf_bytes)))
        writer.encrypt("secret")
        locked = io.BytesIO()
        writer.write(locked)
        with self.assertRaises(HTTPException) as ctx:
            validate_pdf_bytes(locked.getvalue())
        self.assertEqual(ctx.exception.status_code, 400)

    def test_upload_sanitises_filename_in_storage_path(self):
        res = self.upload(self.pdf_bytes, filename="..\\..\\evil name.pdf")
        self.assertEqual(res.status_code, 200)
        body = res.json()
        self.assertEqual(body["filename"], "evil_name.pdf")

        doc = self.rows("documents")[0]
        self.assertEqual(doc["file_path"], f"{settings.DEFAULT_USER_ID}/{body['document_id']}/evil_name.pdf")
        self.assertEqual(self.background.call_args.kwargs["filename"], "evil_name.pdf")
        self.assertEqual(self.background.call_args.kwargs["kind"], "pdf")

    def test_upload_rejects_scanned_pdf_with_clear_message(self):
        with open(os.path.join(self.fixtures_dir.name, "scanned.pdf"), "rb") as f,                 patch.object(documents, "get_ocr_provider", return_value=None):
            res = self.upload(f.read())
        self.assertEqual(res.status_code, 422)
        self.assertIn("scan", res.json()["detail"])
        self.assertEqual(self.rows("documents"), [])

    def test_upload_accepts_docx_and_rejects_fake_docx(self):
        from docx import Document
        buffer = io.BytesIO()
        doc = Document()
        doc.add_heading("Chapter One", level=1)
        doc.add_paragraph("Body text.")
        doc.save(buffer)

        res = self.upload(buffer.getvalue(), filename="Bài giảng.docx")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["kind"], "docx")
        self.assertEqual(self.background.call_args.kwargs["kind"], "docx")
        self.assertTrue(self.rows("documents")[0]["file_path"].endswith(".docx"))

        self.assertEqual(self.upload(b"PK" + bytes([3, 4]) + b" not a real docx", filename="fake.docx").status_code, 400)
        self.assertEqual(self.upload(self.pdf_bytes, filename="renamed.docx").status_code, 400)


class PrivateStorageTests(LocalStoreTestCase):
    def store_pdf(self, user_id=None):
        """Upload thật qua API; chỉ thay bước chunk/embedding/AI bằng mock."""
        with patch.object(document_pipeline, "index_document", Mock()), \
                patch.object(settings, "DEFAULT_USER_ID", user_id or settings.DEFAULT_USER_ID):
            res = self.upload(self.pdf_bytes, filename="my book.pdf")
        self.assertEqual(res.status_code, 200)
        return res.json()["document_id"]

    def test_file_url_serves_original_pdf_through_signed_link(self):
        doc_id = self.store_pdf()
        stored = os.path.join(self.storage_dir, "documents", settings.DEFAULT_USER_ID, doc_id, "my_book.pdf")
        self.assertTrue(os.path.isfile(stored))

        res = self.client.get(f"/api/v1/documents/{doc_id}/file-url")
        self.assertEqual(res.status_code, 200)
        body = res.json()
        self.assertEqual(body["expires_in"], settings.SIGNED_URL_TTL_SECONDS)
        self.assertTrue(body["url"].startswith("http://testserver/local-storage/"))

        download = self.client.get(body["url"])
        self.assertEqual(download.status_code, 200)
        self.assertEqual(download.content, self.pdf_bytes)

    def test_signed_link_rejects_tampered_and_expired_tokens(self):
        doc_id = self.store_pdf()
        url = self.client.get(f"/api/v1/documents/{doc_id}/file-url").json()["url"]
        self.assertEqual(self.client.get(url[:-4] + "AAAA").status_code, 404)
        self.assertEqual(self.client.get("/local-storage/not-a-token").status_code, 404)

        with patch.object(settings, "SIGNED_URL_TTL_SECONDS", -5):
            expired = self.client.get(f"/api/v1/documents/{doc_id}/file-url").json()["url"]
        self.assertEqual(self.client.get(expired).status_code, 404)

    def test_file_url_hidden_from_other_users_and_missing_files(self):
        foreign_doc = self.store_pdf(user_id=OTHER_USER)
        self.assertEqual(self.client.get(f"/api/v1/documents/{foreign_doc}/file-url").status_code, 404)
        self.assertEqual(self.client.get("/api/v1/documents/does-not-exist/file-url").status_code, 404)

        self.store.table("documents").insert({
            "id": "text-doc", "user_id": settings.DEFAULT_USER_ID, "title": "t", "source_type": "text",
        }).execute()
        self.assertEqual(self.client.get("/api/v1/documents/text-doc/file-url").status_code, 404)

    def test_local_storage_blocks_path_traversal(self):
        bucket = self.store.storage.from_("documents")
        for bad in ("../outside.pdf", "a/../../outside.pdf", "..\\outside.pdf", ""):
            with self.assertRaises(ValueError):
                bucket.upload(bad, b"x")
        self.assertFalse(os.path.exists(os.path.join(self.storage_dir, "outside.pdf")))


class DocumentPipelineTests(LocalStoreTestCase):
    def test_auto_generated_flashcards_get_review_state(self):
        topic = {"name": "Memory", "slug": "memory", "description": "d", "difficulty_level": 1}
        card = {"front": "Q", "back": "A", "source_chunk": "s", "card_type": "concept"}
        text = "Memory improves with spaced practice. " * 40
        with patch.object(document_pipeline.RAGEngine, "generate_embeddings",
                          side_effect=lambda texts: [[0.0] * 4 for _ in texts]), \
                patch.object(document_pipeline.AIEngine, "extract_topics_and_graph",
                             return_value={"topics": [topic], "relationships": []}), \
                patch.object(document_pipeline.AIEngine, "generate_flashcards", return_value=[card]):
            document_pipeline.index_document(
                doc_id="doc-1", user_id=settings.DEFAULT_USER_ID, title="t", raw_content=text, supabase=self.store,
            )

        self.assertGreater(len(self.rows("document_chunks")), 0)
        self.assertEqual(len(self.rows("topics")), 1)
        cards = self.rows("flashcards")
        states = self.rows("flashcard_states")
        self.assertEqual(len(cards), 1)
        self.assertEqual([s["flashcard_id"] for s in states], [cards[0]["id"]])
        self.assertEqual(states[0]["user_id"], settings.DEFAULT_USER_ID)
        self.assertTrue(states[0]["next_review_at"])


class FixtureTests(LocalStoreTestCase):
    def read(self, name):
        return PdfReader(os.path.join(self.fixtures_dir.name, f"{name}.pdf"))

    def test_fixture_files_and_expected_json_are_created(self):
        for name in ("text_book", "scanned", "mini_ielts"):
            reader = self.read(name)
            with open(os.path.join(self.fixtures_dir.name, f"{name}.expected.json"), encoding="utf-8") as f:
                self.assertEqual(json.load(f)["page_count"], len(reader.pages))

    def test_text_book_has_six_pages_with_repeated_header_and_footer(self):
        pages = [p.extract_text() for p in self.read("text_book").pages]
        self.assertEqual(len(pages), 6)
        for number, text in enumerate(pages, start=1):
            self.assertIn(self.expected["text_book"]["header"], text)
            self.assertIn(f"Page {number}", text)
        for chapter in self.expected["text_book"]["chapters"]:
            self.assertIn(chapter["title"], pages[chapter["page"] - 1])

    def test_scanned_pdf_has_no_text_layer(self):
        reader = self.read("scanned")
        self.assertEqual(len(reader.pages), 2)
        self.assertEqual("".join((p.extract_text() or "") for p in reader.pages).strip(), "")

    def test_mini_ielts_contains_a_complete_test(self):
        text = "\n".join(p.extract_text() for p in self.read("mini_ielts").pages)
        for marker in ("Test 1", "LISTENING", "PART 4", "READING PASSAGE 3", "WRITING TASK 1", "WRITING TASK 2",
                       "SPEAKING", "Audioscripts", "Listening and Reading answer keys", "21&22   IN EITHER ORDER"):
            self.assertIn(marker, text)

        expected = self.expected["mini_ielts"]
        for skill in ("listening", "reading"):
            self.assertEqual([q["number"] for q in expected[skill]["questions"]], list(range(1, 41)))
            self.assertEqual(len(expected[skill]["answer_key"]), 40)
        self.assertEqual(len(expected["reading"]["passages"]), 3)
        self.assertTrue(all(p["word_count"] >= 400 for p in expected["reading"]["passages"]))
        self.assertEqual(expected["reading"]["answer_key"]["40"]["accept"], ["signature", "a signature"])


if __name__ == "__main__":
    unittest.main()
