import io
import os
import time
import unittest
from unittest.mock import Mock, patch

from docx import Document

from app.config import settings
from app.services import document_pipeline, rag_engine
from app.services.doc_ingest import blocks_from_text, blocks_to_text, extract_blocks
from test_uploads import OTHER_USER, LocalStoreTestCase


def build_docx() -> bytes:
    doc = Document()
    doc.add_heading("Unit 1: Greetings", level=1)
    paragraph = doc.add_paragraph("Say ")
    paragraph.add_run("hello").bold = True
    paragraph.add_run(" when you arrive.")
    doc.add_heading("Practice", level=2)
    doc.add_paragraph("Repeat the dialogue.", style="List Bullet")
    doc.add_paragraph("Swap roles.", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    for (r, c), text in {(0, 0): "Word", (0, 1): "Meaning", (1, 0): "hello", (1, 1): "xin chào"}.items():
        table.cell(r, c).text = text
    doc.add_paragraph("")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class ExtractionTests(LocalStoreTestCase):
    def blocks(self, name="text_book"):
        with open(os.path.join(self.fixtures_dir.name, f"{name}.pdf"), "rb") as f:
            return extract_blocks(f.read(), "pdf")

    def test_pdf_headings_follow_font_size(self):
        result = self.blocks()
        expected = self.expected["text_book"]
        self.assertEqual(result.page_count, 6)
        self.assertEqual(result.empty_pages, 0)
        headings = [(b["level"], b["text"], b["page_number"]) for b in result.blocks if b["block_type"] == "heading"]
        self.assertEqual([h for h in headings if h[0] == 1],
                         [(1, c["title"], c["page"]) for c in expected["chapters"]])
        self.assertEqual([h[1] for h in headings if h[0] == 2], expected["sections"])

    def test_pdf_drops_running_header_and_page_numbers(self):
        text = blocks_to_text(self.blocks().blocks)
        self.assertNotIn(self.expected["text_book"]["header"], text)
        self.assertNotRegex(text, r"Page \d")

    def test_pdf_joins_hyphenated_words_and_wrapped_lines(self):
        text = blocks_to_text(self.blocks().blocks)
        self.assertIn("a method in which learning sessions are separated", text)
        self.assertNotIn("learn-", text)
        # một đoạn = một khối, không gãy theo dòng của PDF
        self.assertTrue(any(b["text"].startswith("New information enters memory") and b["text"].endswith("already knows.")
                            for b in self.blocks().blocks))

    def test_pdf_list_items_and_two_column_order(self):
        blocks = self.blocks().blocks
        items = [b["text"] for b in blocks if b["block_type"] == "list_item"]
        self.assertEqual(items, [f"- {item}" for item in self.expected["text_book"]["list_items"]["items"]])
        page6 = [b["text"].split()[0] for b in blocks if b["page_number"] == 6 and b["block_type"] == "paragraph"]
        self.assertEqual(page6, self.expected["text_book"]["two_column"]["reading_order"])

    def test_scanned_pdf_reports_pages_without_text(self):
        result = self.blocks("scanned")
        self.assertEqual((result.page_count, result.empty_pages, result.blocks), (2, 2, []))

    def test_docx_uses_word_styles_bold_and_tables(self):
        blocks = extract_blocks(build_docx(), "docx").blocks
        self.assertEqual([(b["block_type"], b["level"], b["text"]) for b in blocks], [
            ("heading", 1, "Unit 1: Greetings"),
            ("paragraph", None, "Say **hello** when you arrive."),
            ("heading", 2, "Practice"),
            ("list_item", None, "Repeat the dialogue."),
            ("list_item", None, "Swap roles."),
            ("table", None, "Word | Meaning\nhello | xin chào"),
        ])

    def test_plain_text_becomes_paragraph_blocks(self):
        blocks = blocks_from_text("# Title\n\nFirst line\nsecond line.\n\n- a\n- b\n\n\n")
        self.assertEqual([(b["block_type"], b["text"]) for b in blocks], [
            ("heading", "Title"), ("paragraph", "First line second line."), ("list_item", "- a"), ("list_item", "- b"),
        ])

    def test_large_pdf_extraction_is_fast(self):
        with open(os.path.join(self.fixtures_dir.name, "mini_ielts.pdf"), "rb") as f:
            data = f.read()
        started = time.perf_counter()
        result = extract_blocks(data, "pdf")
        self.assertEqual(result.page_count, 19)
        self.assertLess(time.perf_counter() - started, 2.0)


class BatchEmbeddingTests(unittest.TestCase):
    def test_embeddings_are_requested_in_batches(self):
        genai = Mock()
        genai.embed_content.side_effect = lambda model, content, task_type: {"embedding": [[0.5, 0.25]] * len(content)}
        with patch.object(rag_engine, "_get_gemini_embedding_model", return_value=(genai, "m")):
            vectors = rag_engine.RAGEngine.generate_embeddings([f"chunk {i}" for i in range(230)])
        self.assertEqual(len(vectors), 230)
        self.assertEqual(genai.embed_content.call_count, 3)  # 100 + 100 + 30 thay vì 230 lời gọi
        self.assertEqual(vectors[0], [0.5, 0.25])

    def test_failed_batch_falls_back_without_losing_alignment(self):
        genai = Mock()
        genai.embed_content.side_effect = RuntimeError("quota")
        with patch.object(rag_engine, "_get_gemini_embedding_model", return_value=(genai, "m")):
            vectors = rag_engine.RAGEngine.generate_embeddings(["a", "b", "c"])
        self.assertEqual([len(v) for v in vectors], [768, 768, 768])


class LocalStoreConcurrencyTests(LocalStoreTestCase):
    def test_parallel_writes_do_not_corrupt_or_lose_rows(self):
        from concurrent.futures import ThreadPoolExecutor

        def write(i):
            self.store.table("document_notes").insert({"id": f"n{i}", "content": "x" * 2000}).execute()
            self.store.table("document_notes").update({"content": f"edited {i}"}).eq("id", f"n{i}").execute()

        with ThreadPoolExecutor(max_workers=8) as pool:
            list(pool.map(write, range(40)))
        rows = self.rows("document_notes")
        self.assertEqual(sorted(r["id"] for r in rows), sorted(f"n{i}" for i in range(40)))
        self.assertTrue(all(r["content"].startswith("edited") for r in rows))


class EbookApiTestCase(LocalStoreTestCase):
    def setUp(self):
        super().setUp()
        self.index = Mock()
        p = patch.object(document_pipeline, "index_document", self.index)
        p.start()
        self.addCleanup(p.stop)

    def upload_book(self, user_id=None):
        with patch.object(settings, "DEFAULT_USER_ID", user_id or settings.DEFAULT_USER_ID):
            res = self.upload(self.pdf_bytes)
        self.assertEqual(res.status_code, 200)
        return res.json()["document_id"]

    def blocks(self, doc_id, **params):
        res = self.client.get(f"/api/v1/documents/{doc_id}/blocks", params=params)
        self.assertEqual(res.status_code, 200)
        return res.json()

    def all_blocks(self, doc_id):
        return self.blocks(doc_id, limit=500)["blocks"]


class EbookReadTests(EbookApiTestCase):
    def test_upload_makes_book_readable_before_ai_indexing(self):
        doc_id = self.upload_book()
        self.index.assert_called_once()  # bước AI được gọi sau khi khối đã lưu
        self.assertEqual(len(self.rows("document_blocks")), 32)
        status = self.client.get(f"/api/v1/documents/{doc_id}/status").json()
        self.assertTrue(status["readable"])
        self.assertIn("Chapter 1: How Memory Works", self.index.call_args.args[3])

    def test_blocks_are_paginated_in_reading_order(self):
        doc_id = self.upload_book()
        first = self.blocks(doc_id, limit=10)
        self.assertEqual(len(first["blocks"]), 10)
        self.assertTrue(first["has_more"])
        self.assertEqual(first["blocks"][0]["text"], "Chapter 1: How Memory Works")

        seen, after = [], -1
        while True:
            page = self.blocks(doc_id, limit=10, after=after)
            seen += page["blocks"]
            if not page["has_more"]:
                break
            after = page["next_after"]
        self.assertEqual(len(seen), 32)
        orders = [b["order_index"] for b in seen]
        self.assertEqual(orders, sorted(orders))
        self.assertNotIn("original_text", seen[0])

    def test_outline_lists_headings(self):
        doc_id = self.upload_book()
        outline = self.client.get(f"/api/v1/documents/{doc_id}/outline").json()["outline"]
        self.assertEqual([o["text"] for o in outline if o["level"] == 1],
                         [c["title"] for c in self.expected["text_book"]["chapters"]])
        self.assertEqual(len(outline), 10)
        self.assertEqual(outline[-1]["page_number"], 6)

    def test_legacy_document_gets_blocks_on_first_open(self):
        self.store.table("documents").insert({
            "id": "legacy", "user_id": settings.DEFAULT_USER_ID, "title": "Old", "source_type": "text",
            "raw_content": "First paragraph.\n\nSecond paragraph.", "processing_status": "completed",
        }).execute()
        self.assertEqual([b["text"] for b in self.all_blocks("legacy")], ["First paragraph.", "Second paragraph."])
        self.assertEqual(len(self.all_blocks("legacy")), 2)  # không dựng lại lần hai

    def test_text_upload_is_readable(self):
        res = self.client.post("/api/v1/documents/upload", json={"title": "Note", "raw_content": "One.\n\nTwo."})
        self.assertEqual([b["text"] for b in self.all_blocks(res.json()["document_id"])], ["One.", "Two."])

    def test_other_users_book_is_hidden(self):
        foreign = self.upload_book(user_id=OTHER_USER)
        block_id = self.rows("document_blocks")[0]["id"]
        calls = [
            ("get", f"/api/v1/documents/{foreign}/blocks", None),
            ("get", f"/api/v1/documents/{foreign}/outline", None),
            ("get", f"/api/v1/documents/{foreign}/notes", None),
            ("patch", f"/api/v1/documents/{foreign}/blocks/{block_id}", {"text": "hacked"}),
            ("delete", f"/api/v1/documents/{foreign}/blocks/{block_id}", None),
            ("post", f"/api/v1/documents/{foreign}/notes", {"content": "x"}),
        ]
        for method, url, body in calls:
            res = self.client.request(method, url, json=body)
            self.assertEqual(res.status_code, 404, url)
        self.assertNotEqual(self.rows("document_blocks")[0]["text"], "hacked")


class EbookEditTests(EbookApiTestCase):
    def setUp(self):
        super().setUp()
        self.doc_id = self.upload_book()
        self.blocks_before = self.all_blocks(self.doc_id)
        self.url = f"/api/v1/documents/{self.doc_id}/blocks"

    def test_edit_keeps_original_and_can_be_restored(self):
        block = self.blocks_before[2]
        res = self.client.patch(f"{self.url}/{block['id']}", json={"text": "Edited with **bold** word."})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["text"], "Edited with **bold** word.")
        self.assertTrue(res.json()["is_edited"])
        self.assertEqual(self.all_blocks(self.doc_id)[2]["text"], "Edited with **bold** word.")

        stored = [r for r in self.rows("document_blocks") if r["id"] == block["id"]][0]
        self.assertEqual(stored["original_text"], block["text"])

        restored = self.client.post(f"{self.url}/{block['id']}/restore").json()
        self.assertEqual(restored["text"], block["text"])
        self.assertFalse(restored["is_edited"])

    def test_change_block_type_sets_and_clears_level(self):
        block = self.blocks_before[2]
        heading = self.client.patch(f"{self.url}/{block['id']}", json={"block_type": "heading", "level": 3}).json()
        self.assertEqual((heading["block_type"], heading["level"]), ("heading", 3))
        self.assertEqual(len(self.client.get(f"/api/v1/documents/{self.doc_id}/outline").json()["outline"]), 11)

        paragraph = self.client.patch(f"{self.url}/{block['id']}", json={"block_type": "paragraph"}).json()
        self.assertEqual((paragraph["block_type"], paragraph["level"]), ("paragraph", None))
        self.assertEqual(self.client.patch(f"{self.url}/{block['id']}", json={"block_type": "video"}).status_code, 422)

    def test_insert_block_lands_between_neighbours(self):
        anchor = self.blocks_before[4]
        res = self.client.post(self.url, json={"after_block_id": anchor["id"], "text": "My own paragraph."})
        self.assertEqual(res.status_code, 200)
        texts = [b["text"] for b in self.all_blocks(self.doc_id)]
        self.assertEqual(texts[4:7], [anchor["text"], "My own paragraph.", self.blocks_before[5]["text"]])

        self.client.post(self.url, json={"after_block_id": self.blocks_before[-1]["id"], "text": "The end."})
        self.client.post(self.url, json={"text": "Preface.", "block_type": "heading"})
        after = self.all_blocks(self.doc_id)
        self.assertEqual((after[0]["text"], after[0]["level"], after[-1]["text"]), ("Preface.", 2, "The end."))
        self.assertEqual(len(after), 35)

        added = res.json()["id"]
        self.assertEqual(self.client.post(f"{self.url}/{added}/restore").status_code, 409)

    def test_delete_block(self):
        target = self.blocks_before[3]
        self.assertEqual(self.client.delete(f"{self.url}/{target['id']}").status_code, 200)
        remaining = self.all_blocks(self.doc_id)
        self.assertEqual(len(remaining), 31)
        self.assertNotIn(target["id"], [b["id"] for b in remaining])
        self.assertEqual(self.client.delete(f"{self.url}/{target['id']}").status_code, 404)


class NotebookTests(EbookApiTestCase):
    def test_notes_crud_scoped_to_book(self):
        doc_id = self.upload_book()
        other_doc = self.upload_book()
        url = f"/api/v1/documents/{doc_id}/notes"
        block = self.all_blocks(doc_id)[2]

        free = self.client.post(url, json={"content": "  Ôn lại chương 1 trước thứ Sáu.  "})
        self.assertEqual(free.status_code, 200)
        self.assertEqual(free.json()["content"], "Ôn lại chương 1 trước thứ Sáu.")
        quoted = self.client.post(url, json={"content": "Ý chính", "quote": "attention", "block_id": block["id"]}).json()
        self.assertEqual((quoted["quote"], quoted["block_id"]), ("attention", block["id"]))

        notes = self.client.get(url).json()["notes"]
        self.assertEqual({n["id"] for n in notes}, {free.json()["id"], quoted["id"]})
        self.assertEqual(self.client.get(f"/api/v1/documents/{other_doc}/notes").json()["notes"], [])

        edited = self.client.patch(f"{url}/{quoted['id']}", json={"content": "Ý chính (đã sửa)"})
        self.assertEqual(edited.json()["content"], "Ý chính (đã sửa)")
        self.assertEqual(edited.json()["quote"], "attention")

        self.assertEqual(self.client.delete(f"{url}/{free.json()['id']}").status_code, 200)
        self.assertEqual([n["id"] for n in self.client.get(url).json()["notes"]], [quoted["id"]])

    def test_note_validation(self):
        doc_id = self.upload_book()
        url = f"/api/v1/documents/{doc_id}/notes"
        self.assertEqual(self.client.post(url, json={"content": ""}).status_code, 422)
        self.assertEqual(self.client.post(url, json={"content": "x", "block_id": "missing"}).status_code, 404)
        self.assertEqual(self.client.patch(f"{url}/missing", json={"content": "x"}).status_code, 404)
        # ghi chú của sách khác không sửa được qua URL của sách này
        other_note = self.client.post(f"/api/v1/documents/{self.upload_book()}/notes", json={"content": "a"}).json()
        self.assertEqual(self.client.patch(f"{url}/{other_note['id']}", json={"content": "b"}).status_code, 404)


if __name__ == "__main__":
    unittest.main()
